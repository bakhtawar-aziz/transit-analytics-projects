"""Compute reliability KPIs from SQLite and export CSV + Excel (with native charts) for Power BI/Excel."""
import sqlite3, pandas as pd, pathlib
from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, Reference
ROOT = pathlib.Path(__file__).resolve().parents[1]; OUT = ROOT / "output"
con = sqlite3.connect(OUT / "fleet.db"); q = lambda s: pd.read_sql(s, con)
monthly = q("SELECT * FROM v_monthly_reliability ORDER BY vehicle_type, month")
by_type = q("""SELECT vehicle_type, ROUND(SUM(km),0) km, SUM(failures) failures, SUM(road_calls) road_calls,
   ROUND(SUM(failures)*100000.0/SUM(km),2) failures_per_100k_km, ROUND(SUM(km)*1.0/SUM(failures),0) mdbf_km,
   ROUND(SUM(road_calls)*100000.0/SUM(km),2) road_calls_per_100k_km FROM v_monthly_reliability GROUP BY vehicle_type""")
avail = q("""SELECT division, SUM(vehicle_days) vehicle_days, SUM(available_days) available_days,
   ROUND(100.0*SUM(available_days)/SUM(vehicle_days),2) availability_pct FROM v_availability GROUP BY division""")
avail_m = q("SELECT * FROM v_availability ORDER BY division, month")
by_system = q("""SELECT system, COUNT(*) failures, ROUND(100.0*COUNT(*)/(SELECT COUNT(*) FROM failures),1) pct
   FROM failures GROUP BY system ORDER BY failures DESC""")
age = q("""SELECT CASE WHEN 2025-v.model_year<5 THEN '0-4 yrs' WHEN 2025-v.model_year<10 THEN '5-9 yrs' ELSE '10+ yrs' END age_band,
   COUNT(DISTINCT v.vehicle_id) vehicles, SUM(d.km) km FROM vehicles v JOIN daily_mileage d USING(vehicle_id) GROUP BY 1""")
agef = q("""SELECT CASE WHEN 2025-v.model_year<5 THEN '0-4 yrs' WHEN 2025-v.model_year<10 THEN '5-9 yrs' ELSE '10+ yrs' END age_band,
   COUNT(*) failures FROM failures f JOIN vehicles v USING(vehicle_id) GROUP BY 1""")
age = age.merge(agef, on="age_band"); age["failures_per_100k_km"] = (age.failures * 1e5 / age.km).round(2)
trace = q("SELECT status, COUNT(*) n FROM v_traceability GROUP BY status")
wo = q("""SELECT ROUND(AVG(labour_hours),1) avg_labour_h, ROUND(AVG(julianday(closed_date)-julianday(opened_date)),1) avg_days_open,
   ROUND(SUM(parts_cost),0) parts_cost FROM work_orders WHERE orphan_flag=0""")
tabs = dict(KPI_ByType=by_type, KPI_Monthly=monthly, KPI_Availability=avail, Availability_Monthly=avail_m,
            FailuresBySystem=by_system, AgeBands=age, Traceability=trace, WorkOrderStats=wo,
            ErrorLog=pd.read_csv(OUT / "error_log.csv"), MileageAudit=pd.read_csv(OUT / "mileage_audit.csv"))
for n, d in tabs.items(): d.to_csv(OUT / f"{n}.csv", index=False)
xl = OUT / "fleet_reliability_kpis.xlsx"
with pd.ExcelWriter(xl) as w:
    pd.DataFrame({"note": ["SYNTHETIC DATA - generated for portfolio use; not real agency data.",
        "Failures per 100,000 km = failures / km * 100000. MDBF = km / failures. Availability = non-OOS vehicle-days / all vehicle-days."]}).to_excel(w, sheet_name="README", index=False)
    for n, d in tabs.items(): d.to_excel(w, sheet_name=n[:31], index=False)
wb = load_workbook(xl)
ws = wb["KPI_ByType"]; ch = BarChart(); ch.title = "Failures per 100,000 km by vehicle type"
ch.add_data(Reference(ws, min_col=5, min_row=1, max_row=ws.max_row), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)); ws.add_chart(ch, "J2")
wm = wb["KPI_Monthly"]; lc = LineChart(); lc.title = "Monthly failures / 100k km (all types stacked by row)"
lc.add_data(Reference(wm, min_col=6, min_row=1, max_row=wm.max_row), titles_from_data=True); wm.add_chart(lc, "J2")
for s in wb.worksheets:
    for col in s.columns: s.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col[:50]) + 2))
wb.save(xl); print("KPI tables + workbook written"); print(by_type.to_string(index=False)); print(avail.to_string(index=False)); print(age.to_string(index=False))
